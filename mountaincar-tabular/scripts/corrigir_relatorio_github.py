"""Correção editorial estrita do PDF acadêmico e geração do DOCX editável.

Preserva resultados científicos, páginas e figuras. Falha sem sobrescrever
o PDF caso a estrutura fonte não coincida com o relatório auditado.
"""
from pathlib import Path
import fitz

BASE = Path("mountaincar-tabular/entrega_final")
REPORT = BASE / "CMP629_MOUNTAINCAR_RELATORIO_DEFINITIVO.pdf"
WORD = BASE / "CMP629_MOUNTAINCAR_RELATORIO_EDITAVEL.docx"
NAMES = [
    "Caetano gubert meneghetti",
    "Fernando bosquetti tedesco",
    "Gabriel de senne amorim",
    "José antônio coutinho de moraes filho",
]

def lines(page):
    for block in page.get_text("dict")["blocks"]:
        if block["type"] != 0:
            continue
        for line in block["lines"]:
            tx = "".join(sp["text"] for sp in line["spans"]).strip()
            if tx:
                yield tx, fitz.Rect(line["bbox"])

def replace_line(page, matcher, new_text, size=11.4):
    matches = [rect for text, rect in lines(page) if matcher(text)]
    if len(matches) != 1:
        raise AssertionError(f"Linha unica esperada ({len(matches)}): {new_text}")
    rect = matches[0]
    page.add_redact_annot(rect + (-1, -1, 2, 1), fill=(1, 1, 1))
    page.apply_redactions(images=0, graphics=0, text=0)
    page.insert_text((rect.x0, rect.y1 - 1), new_text,
                     fontname="helv", fontsize=size, color=(0, 0, 0))

source = fitz.open(REPORT)
original = [page.get_text() for page in source]
assert len(source) == 29, f"Esperadas 29 paginas; encontradas {len(source)}"

cover = source[0]
assert all(any(key in text.upper() for text, _ in lines(cover))
           for key in ["GUBERT", "BOSQUETTI", "SENNE", "MORAES FILHO"])
cover.add_redact_annot(fitz.Rect(70.8, 171.0, 541.0, 231.0), fill=(1, 1, 1))
cover.apply_redactions(images=0, graphics=0, text=0)
for name, y in zip(NAMES, [183.6, 198.1, 212.5, 227.0]):
    cover.insert_text((72, y), name, fontsize=11.4,
                      fontname="helv", color=(0, 0, 0))

# Evita a citação órfã ao enunciado formalmente excluído das referências.
replace_line(
    cover,
    lambda text: text.startswith("(Universidade Federal do Rio Grande do Sul, s.d.)"),
    "e os fundamentos de Sutton e Barto (2018, caps. 6, 8 e 12)."
)
continuation = [rect for text, rect in lines(cover)
                if text.strip() == "(2018, caps. 6, 8 e 12)."]
assert len(continuation) == 1
cover.add_redact_annot(continuation[0] + (-1, -1, 2, 1), fill=(1, 1, 1))
cover.apply_redactions(images=0, graphics=0, text=0)

# Página 25 contém a seção REFERÊNCIAS: manter somente Sutton e Barto.
references = source[24]
source_refs = references.get_text().upper()
assert "ANTONIOMORAES85." in source_refs
assert "UNIVERSIDADE FEDERAL DO RIO GRANDE DO SUL." in source_refs
assert "SUTTON, RICHARD S.; BARTO, ANDREW G." in source_refs
references.add_redact_annot(fitz.Rect(70.5, 100, 541, 306),
                           fill=(1, 1, 1))
references.apply_redactions(images=0, graphics=0, text=0)
references.insert_text(
    (72, 117.8),
    "SUTTON, Richard S.; BARTO, Andrew G. Reinforcement learning: an introduction. 2.",
    fontsize=11.45, fontname="helv"
)
references.insert_text(
    (72, 133.3),
    "ed. Cambridge, MA: MIT Press, 2018.",
    fontsize=11.45, fontname="helv"
)

# Citação ao repositório no apêndice permanece como proveniência, sem autor/ano.
appendix = source[25]
matches = [(text, rect) for text, rect in lines(appendix)
           if text.lower().startswith("antoniomoraes85 (2026).")]
assert len(matches) == 1
old_text = matches[0][0]
replace_line(
    appendix,
    lambda text: text == old_text,
    old_text.replace("antoniomoraes85 (2026).", "de referência.", 1),
    size=10.85
)

assert all(name in cover.get_text() for name in NAMES)
assert "ANTONIOMORAES85." not in references.get_text().upper()
assert "UNIVERSIDADE FEDERAL DO RIO GRANDE DO SUL." not in references.get_text().upper()
assert "SUTTON, Richard S." in references.get_text()
assert "Universidade Federal do Rio Grande do Sul, s.d." not in cover.get_text()
assert "antoniomoraes85 (2026)." not in appendix.get_text()
for index in range(len(source)):
    if index not in [0, 24, 25]:
        assert original[index] == source[index].get_text(), (
            f"Pagina indevidamente alterada: {index + 1}"
        )
source.set_metadata({
    **source.metadata,
    "author": "; ".join(NAMES),
    "subject": "CMP629 MountainCar - correção editorial de 13/10/2026"
})
temp = REPORT.with_name("CMP629_MOUNTAINCAR_TEMP.pdf")
source.save(temp, garbage=4, deflate=True)
source.close()
assert fitz.open(temp).page_count == 29
temp.replace(REPORT)
print("PDF_29_PAGINAS_CORRIGIDO", REPORT, REPORT.stat().st_size)

from pdf2docx import Converter
converter = Converter(str(REPORT))
try:
    converter.convert(str(WORD), start=0, end=None)
finally:
    converter.close()
assert WORD.exists() and WORD.stat().st_size > 100000
from docx import Document
word = Document(WORD)
all_text = "\n".join(paragraph.text for paragraph in word.paragraphs)
all_text += "\n" + "\n".join(
    cell.text for table in word.tables for row in table.rows for cell in row.cells
)
assert all(name in all_text for name in NAMES), "Autoria incompleta no Word."
assert "ANTONIOMORAES85." not in all_text
assert "UNIVERSIDADE FEDERAL DO RIO GRANDE DO SUL." not in all_text
print("DOCX_29_PAGINAS_CONVERTIDO", WORD, WORD.stat().st_size,
      "tabelas", len(word.tables))
