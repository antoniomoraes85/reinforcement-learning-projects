"""Empacota apenas os tres entregaveis revisados, com manifesto e verificacoes."""
import hashlib, json, zipfile
from pathlib import Path
import fitz
from docx import Document

base=Path("mountaincar-tabular/entrega_final")
names=[
    "Caetano gubert meneghetti",
    "Fernando bosquetti tedesco",
    "Gabriel de senne amorim",
    "José antônio coutinho de moraes filho",
]
paths=[
    base/"CMP629_MOUNTAINCAR_ENTREGA_FINAL.ipynb",
    base/"CMP629_MOUNTAINCAR_RELATORIO_DEFINITIVO.pdf",
    base/"CMP629_MOUNTAINCAR_RELATORIO_EDITAVEL.docx",
]
assert all(p.exists() for p in paths)
nb=json.loads(paths[0].read_text(encoding="utf-8"))
assert len(nb["cells"])==31
header="".join(nb["cells"][0]["source"])
positions=[header.find(name) for name in names]
assert positions==sorted(positions) and min(positions)>=0
pdf=fitz.open(paths[1])
assert pdf.page_count==29
assert all(name in pdf[0].get_text() for name in names)
refs=pdf[24].get_text().upper()
assert "SUTTON, RICHARD S." in refs
assert "ANTONIOMORAES85." not in refs
assert "UNIVERSIDADE FEDERAL DO RIO GRANDE DO SUL." not in refs
word=Document(paths[2])
tx="\n".join(p.text for p in word.paragraphs)
tx+="\n"+"\n".join(c.text for t in word.tables for r in t.rows for c in r.cells)
assert all(name in tx for name in names)
info=base/"LEIA_ME_GRUPO.md"
info.write_text("""# CMP629 - Entrega final MountainCar (grupo)

## Arquivos deste pacote
- CMP629_MOUNTAINCAR_ENTREGA_FINAL.ipynb - notebook obrigatório; inclui resultados, gráficos, vídeos e código de apresentação. Não executa novos treinamentos ao usar o fluxo padrão.
- CMP629_MOUNTAINCAR_RELATORIO_DEFINITIVO.pdf - relatório científico completo revisado, 29 páginas.
- CMP629_MOUNTAINCAR_RELATORIO_EDITAVEL.docx - versão editável do relatório.

## Revisão 13/10/2026
- Integrantes em ordem alfabética: Caetano gubert meneghetti; Fernando bosquetti tedesco; Gabriel de senne amorim; José antônio coutinho de moraes filho.
- Retiradas do rol bibliográfico as referências ANTONIOMORAES85 (2026) e UNIVERSIDADE FEDERAL DO RIO GRANDE DO SUL (enunciado CMP629).
- Resultados científicos mantidos: 108 execuções exploratórias e 60 finais, com dez seeds por método e grade; avaliações independentes de 100 episódios por política.
- Não foi executado novo treinamento nem nova avaliação de políticas nesta revisão editorial.

## Antes da submissão
1. Cada integrante revisa nomes, interpretação e conteúdo científico, e confirma a autoria.
2. Abrir o notebook no Google Colab hospedado e testar a execução padrão (análise sem treinamento). O teste remoto ainda não foi registrado como concluído.
3. Confirmar com o docente se existe exigência de uso do notebook-base específico.
4. Conferir regras de entrega e eventuais declarações de uso de IA.

Repositório:
https://github.com/antoniomoraes85/reinforcement-learning-projects/tree/main/mountaincar-tabular
""",encoding="utf-8")
manifest={
  "projeto":"CMP629 MountainCar",
  "tipo":"pacote_grupo_revisao_editorial",
  "arquivos":[{"nome":p.name,"tamanho_bytes":p.stat().st_size,
               "sha256":hashlib.sha256(p.read_bytes()).hexdigest()} for p in paths],
  "limitacoes":["colab_hospedado_nao_testado","notebook_base_do_docente_nao_conferido"]
}
out=base/"CMP629_MOUNTAINCAR_PACOTE_GRUPO.zip"
with zipfile.ZipFile(out,"w",compression=zipfile.ZIP_DEFLATED,compresslevel=6,allowZip64=True) as z:
    for p in paths:
        z.write(p,arcname=p.name)
    z.write(info,arcname=info.name)
    z.writestr("MANIFESTO_SHA256.json",json.dumps(manifest,ensure_ascii=False,indent=2))
with zipfile.ZipFile(out) as z:
    assert z.testzip() is None
    assert len(z.namelist())==5
print("PACOTE_GRUPO_VERIFICADO",out,"bytes",out.stat().st_size)
